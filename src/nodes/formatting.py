"""
Node 10: Formatting & Styling

Final document formatting and export to various formats.
"""

from pathlib import Path
from typing import Any, Dict
from datetime import datetime
import uuid

from loguru import logger

from .base import BaseNode
from ..core.state import AgentState, DocumentFormat


class FormattingNode(BaseNode):
    """
    Formats and exports documents to final format.
    
    Supported exports:
    - Markdown
    - PDF (via WeasyPrint)
    - DOCX (via python-docx)
    - HTML
    """
    
    node_name = "formatting"
    max_retries = 2
    
    def validate_input(self, state: AgentState) -> tuple[bool, str]:
        """Check for formatted document."""
        if not state.get("formatted_document"):
            return False, "No formatted document to export"
        return True, ""
    
    def process(self, state: AgentState) -> AgentState:
        """Export document to final format."""
        formatted = state["formatted_document"]
        output_format = state.get("output_format", "markdown")
        
        logger.info(f"Exporting document as: {output_format}")
        
        # Ensure output directory exists
        output_dir = Path(self.config.documents.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        
        # Generate filename
        doc_id = state.get("document_draft", {}).get("id", str(uuid.uuid4())[:8])
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        base_name = f"research_{timestamp}_{doc_id}"
        
        # Export based on format
        try:
            if output_format in ("markdown", "md"):
                output_path = self._export_markdown(formatted, output_dir, base_name)
            elif output_format == "pdf":
                output_path = self._export_pdf(formatted, output_dir, base_name)
            elif output_format == "docx":
                output_path = self._export_docx(formatted, output_dir, base_name)
            elif output_format == "html":
                output_path = self._export_html(formatted, output_dir, base_name)
            else:
                # Default to markdown
                output_path = self._export_markdown(formatted, output_dir, base_name)
            
            state["output_path"] = str(output_path)
            logger.info(f"Document exported to: {output_path}")
            
        except Exception as e:
            logger.error(f"Export failed: {e}")
            self.add_warning(state, f"Export failed: {e}")
            # Try markdown as fallback
            try:
                output_path = self._export_markdown(formatted, output_dir, base_name)
                state["output_path"] = str(output_path)
            except:
                state["output_path"] = ""
        
        return state
    
    def _export_markdown(self, formatted: Dict[str, Any], output_dir: Path, base_name: str) -> Path:
        """Export as Markdown."""
        output_path = output_dir / f"{base_name}.md"
        content = formatted.get("content", "")
        output_path.write_text(content, encoding="utf-8")
        return output_path
    
    def _export_html(self, formatted: Dict[str, Any], output_dir: Path, base_name: str) -> Path:
        """Export as HTML."""
        output_path = output_dir / f"{base_name}.html"
        content = formatted.get("content", "")
        
        # If content is markdown, convert to HTML
        if formatted.get("format") == "markdown":
            try:
                import markdown
                content = markdown.markdown(content, extensions=["tables", "fenced_code"])
                content = f"""<!DOCTYPE html>
<html><head><meta charset="UTF-8"><title>Research Document</title>
<style>body{{font-family:sans-serif;max-width:800px;margin:0 auto;padding:20px;}}</style>
</head><body>{content}</body></html>"""
            except ImportError:
                logger.warning("markdown package not installed")
        
        output_path.write_text(content, encoding="utf-8")
        return output_path
    
    def _export_pdf(self, formatted: Dict[str, Any], output_dir: Path, base_name: str) -> Path:
        """Export as PDF using WeasyPrint."""
        output_path = output_dir / f"{base_name}.pdf"
        content = formatted.get("content", "")
        
        try:
            from weasyprint import HTML, CSS
            
            # Convert content to HTML if needed
            if formatted.get("format") == "markdown":
                try:
                    import markdown
                    html_content = markdown.markdown(content, extensions=["tables", "fenced_code"])
                except ImportError:
                    html_content = f"<pre>{content}</pre>"
            else:
                html_content = content
            
            # Wrap in full HTML document
            full_html = f"""<!DOCTYPE html>
<html><head><meta charset="UTF-8">
<style>
    body {{ font-family: Georgia, serif; font-size: 12pt; line-height: 1.6; max-width: 700px; margin: 0 auto; padding: 40px; }}
    h1 {{ color: #2c3e50; page-break-after: avoid; }}
    h2 {{ color: #34495e; page-break-after: avoid; }}
    pre {{ background: #f5f5f5; padding: 10px; overflow-x: auto; }}
    .references {{ page-break-before: always; }}
</style>
</head><body>{html_content}</body></html>"""
            
            HTML(string=full_html).write_pdf(str(output_path))
            logger.info(f"PDF generated: {output_path}")
            
        except ImportError:
            logger.warning("WeasyPrint not installed, falling back to HTML")
            return self._export_html(formatted, output_dir, base_name)
        except Exception as e:
            logger.error(f"PDF generation failed: {e}")
            return self._export_html(formatted, output_dir, base_name)
        
        return output_path
    
    def _export_docx(self, formatted: Dict[str, Any], output_dir: Path, base_name: str) -> Path:
        """Export as DOCX using python-docx."""
        output_path = output_dir / f"{base_name}.docx"
        content = formatted.get("content", "")
        
        try:
            from docx import Document
            from docx.shared import Pt, Inches
            
            doc = Document()
            
            # Parse content (assuming markdown-ish format)
            lines = content.split("\n")
            
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                
                if line.startswith("# "):
                    doc.add_heading(line[2:], level=0)
                elif line.startswith("## "):
                    doc.add_heading(line[3:], level=1)
                elif line.startswith("### "):
                    doc.add_heading(line[4:], level=2)
                elif line.startswith("- "):
                    doc.add_paragraph(line[2:], style="List Bullet")
                elif line.startswith("---"):
                    doc.add_paragraph("_" * 50)
                else:
                    doc.add_paragraph(line)
            
            doc.save(str(output_path))
            logger.info(f"DOCX generated: {output_path}")
            
        except ImportError:
            logger.warning("python-docx not installed, falling back to markdown")
            return self._export_markdown(formatted, output_dir, base_name)
        except Exception as e:
            logger.error(f"DOCX generation failed: {e}")
            return self._export_markdown(formatted, output_dir, base_name)
        
        return output_path
